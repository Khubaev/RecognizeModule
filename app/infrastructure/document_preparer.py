"""Приведение PDF, JPEG, WebP и Excel к вложению для GigaChat."""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image, UnidentifiedImageError

from app.domain.entities import DocumentFormat, PreparedDocument, SourceDocument
from app.domain.exceptions import InvalidDocumentError


class FileDocumentPreparer:
    """PDF и JPEG передаёт без изменений. WebP конвертирует в JPEG, Excel — в текст."""

    def prepare(self, document: SourceDocument) -> PreparedDocument:
        """Возвращает байты, имя и MIME-тип готового вложения."""
        if document.format is DocumentFormat.PDF:
            return PreparedDocument(document.filename, document.content, "application/pdf")
        if document.format is DocumentFormat.JPEG:
            return PreparedDocument(document.filename, document.content, "image/jpeg")
        if document.format is DocumentFormat.WEBP:
            return self._webp_to_jpeg(document)
        return self._excel_to_text(document)

    def _webp_to_jpeg(self, document: SourceDocument) -> PreparedDocument:
        try:
            image = Image.open(BytesIO(document.content))
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            buffer = BytesIO()
            image.save(buffer, format="JPEG", quality=90)
        except (UnidentifiedImageError, OSError) as exc:
            raise InvalidDocumentError("Не удалось прочитать WebP-изображение.") from exc
        return PreparedDocument(
            filename=f"{Path(document.filename).stem}.jpg",
            content=buffer.getvalue(),
            media_type="image/jpeg",
        )

    def _excel_to_text(self, document: SourceDocument) -> PreparedDocument:
        try:
            workbook = load_workbook(BytesIO(document.content), data_only=True, read_only=True)
        except (InvalidFileException, OSError, ValueError) as exc:
            raise InvalidDocumentError(
                "Не удалось прочитать Excel. Файл .xls сохраните как .xlsx."
            ) from exc
        lines: list[str] = []
        try:
            for sheet in workbook.worksheets:
                lines.append(f"# Лист: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if value is None else str(value) for value in row]
                    if any(cell.strip() for cell in cells):
                        lines.append(" | ".join(cells))
                lines.append("")
        finally:
            workbook.close()
        text = "\n".join(lines).strip() or "(пустая книга)"
        return PreparedDocument(
            filename=f"{Path(document.filename).stem}.txt",
            content=text.encode("utf-8"),
            media_type="text/plain",
        )
